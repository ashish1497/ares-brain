"""Extract plain text from course files: PDF, DOCX, PPTX, XLSX."""
from pathlib import Path


class UnsupportedFormat(Exception):
    pass


def _pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages).strip()


def _docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(parts).strip()


def _pptx(path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append("Notes: " + slide.notes_slide.notes_text_frame.text.strip())
    return "\n".join(out).strip()


def _xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"--- Sheet: {ws.title} ---")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                out.append("\t".join(cells))
    wb.close()
    return "\n".join(out).strip()


_READERS = {".pdf": _pdf, ".docx": _docx, ".pptx": _pptx, ".xlsx": _xlsx}


def extract_text(path: Path) -> str:
    reader = _READERS.get(path.suffix.lower())
    if reader is None:
        raise UnsupportedFormat(f"no extractor for {path.suffix!r} ({path.name})")
    return reader(path)
